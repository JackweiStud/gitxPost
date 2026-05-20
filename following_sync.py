#!/usr/bin/env python3
"""Fetch X following accounts and compare them with radar accounts."""

from __future__ import annotations

import argparse
import json
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

from browser_cdp_session import create_session

try:
    from patchright.sync_api import TimeoutError as PlaywrightTimeoutError
    from patchright.sync_api import Page
except Exception:
    PlaywrightTimeoutError = Exception
    Page = object

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
except Exception:
    Workbook = None

BASE_DIR = Path(__file__).resolve().parent
ACCOUNTS_JSON = BASE_DIR / "xinfo" / "accounts.json"
FOLLOWING_DIR = BASE_DIR / "xinfo" / "log" / "following"
DEFAULT_XLSX = BASE_DIR / "xinfo" / "log" / "myfollowing.xlsx"
DEFAULT_PROFILE_DIR = BASE_DIR / "chrome_data_mirror"

ROW_COLUMNS = [
    "data_label",
    "handle",
    "display_name",
    "note",
    "bio",
    "profile_url",
    "in_following",
    "in_radar",
    "radar_status",
    "language_label",
    "source",
    "fetched_at",
]

SUMMARY_FIELDS = [
    "username",
    "fetched_at",
    "expected_following_count",
    "collected_following_count",
    "completion_status",
    "following_count",
    "radar_active_count",
    "gap1_count",
    "gap2_count",
    "complete",
    "json_snapshot",
    "xlsx_report",
]

SHEET_MAP = [
    ("Following", "following"),
    ("Radar", "radar"),
    ("GAP1_Following_Not_Radar", "gap1"),
    ("GAP2_Radar_Not_Following", "gap2"),
]


def _log(message: str, level: str = "INFO") -> None:
    print(f"[{datetime.now().strftime('%H:%M:%S')}] [{level}] {message}", file=sys.stderr, flush=True)


def normalize_handle(handle: str) -> str:
    return (handle or "").strip().lstrip("@").lower()


def display_handle(handle: str) -> str:
    return (handle or "").strip().lstrip("@")


def _row_from_legacy_user(legacy: dict[str, Any]) -> dict[str, str] | None:
    handle = display_handle(str(legacy.get("screen_name") or ""))
    if not handle:
        return None
    return {
        "handle": handle,
        "display_name": str(legacy.get("name") or handle),
        "bio": str(legacy.get("description") or ""),
        "profile_url": f"https://x.com/{handle}",
    }


def _extract_user_results(value: Any) -> list[dict[str, Any]]:
    results = []
    if isinstance(value, dict):
        user_results = value.get("user_results")
        if isinstance(user_results, dict):
            result = user_results.get("result")
            if isinstance(result, dict):
                results.append(result)
        for child in value.values():
            results.extend(_extract_user_results(child))
    elif isinstance(value, list):
        for child in value:
            results.extend(_extract_user_results(child))
    return results


def extract_user_rows_from_graphql_payload(payload: Any) -> list[dict[str, str]]:
    rows_by_key: dict[str, dict[str, str]] = {}
    for result in _extract_user_results(payload):
        legacy = result.get("legacy")
        if not isinstance(legacy, dict):
            continue
        row = _row_from_legacy_user(legacy)
        if row:
            key = normalize_handle(row["handle"])
            rows_by_key[key] = merge_account_row(rows_by_key.get(key, {}), row)
    return list(rows_by_key.values())


def load_radar_accounts(accounts_json: Path = ACCOUNTS_JSON) -> list[dict[str, Any]]:
    if not accounts_json.exists():
        return []
    data = json.loads(accounts_json.read_text("utf-8"))
    accounts = data.get("accounts", [])
    if not isinstance(accounts, list):
        return []
    return [row for row in accounts if isinstance(row, dict)]


def load_language_by_handle(accounts: list[dict[str, Any]]) -> dict[str, dict[str, Any]]:
    language = {}
    for row in accounts:
        key = normalize_handle(row.get("handle", ""))
        if not key:
            continue
        if row.get("language_label") or row.get("language"):
            language[key] = {
                "language": row.get("language", ""),
                "language_label": row.get("language_label", ""),
            }
    return language


def is_collection_complete(
    *,
    expected_following_count: int | None,
    collected_count: int,
    limit: int | None = 0,
) -> bool:
    return completion_status(
        expected_following_count=expected_following_count,
        collected_count=collected_count,
        limit=limit,
    ) == "complete"


def completion_status(
    *,
    expected_following_count: int | None,
    collected_count: int,
    limit: int | None = 0,
) -> str:
    if expected_following_count is None or expected_following_count < 0:
        return "unknown"
    if collected_count >= expected_following_count:
        return "complete"
    if limit and collected_count >= limit:
        return "partial_limited"
    return "partial"


def _is_better_text(candidate: str, current: str, handle: str = "") -> bool:
    candidate = candidate or ""
    current = current or ""
    if not candidate:
        return False
    if not current:
        return True
    if handle and current.lower() == handle.lower() and candidate.lower() != handle.lower():
        return True
    return len(candidate) > len(current)


def merge_account_row(existing: dict[str, Any], incoming: dict[str, Any]) -> dict[str, Any]:
    if not existing:
        return dict(incoming)
    if not incoming:
        return dict(existing)

    merged = dict(existing)
    handle = display_handle(str(merged.get("handle") or incoming.get("handle") or ""))
    merged["handle"] = merged.get("handle") or incoming.get("handle", "")

    for field in ("display_name", "bio", "profile_url"):
        current = str(merged.get(field) or "")
        candidate = str(incoming.get(field) or "")
        if _is_better_text(candidate, current, handle):
            merged[field] = candidate

    for field, value in incoming.items():
        if field not in merged or merged.get(field) in ("", None):
            merged[field] = value
    return merged


def _following_row(
    item: dict[str, Any],
    *,
    data_label: str,
    in_radar: bool,
    radar: dict[str, Any] | None,
    language_by_handle: dict[str, dict[str, Any]],
    fetched_at: str,
) -> dict[str, Any]:
    handle = display_handle(item.get("handle", ""))
    key = normalize_handle(handle)
    language = language_by_handle.get(key, {})
    return {
        "data_label": data_label,
        "handle": handle,
        "display_name": item.get("display_name", ""),
        "note": (radar or {}).get("note", ""),
        "bio": item.get("bio", ""),
        "profile_url": item.get("profile_url") or f"https://x.com/{handle}",
        "in_following": True,
        "in_radar": in_radar,
        "radar_status": (radar or {}).get("status", ""),
        "language_label": language.get("language_label", ""),
        "source": "x_following",
        "fetched_at": fetched_at,
    }


def _radar_row(
    item: dict[str, Any],
    *,
    data_label: str,
    in_following: bool,
    following: dict[str, Any] | None,
    language_by_handle: dict[str, dict[str, Any]],
    fetched_at: str,
) -> dict[str, Any]:
    handle = display_handle(item.get("handle", ""))
    key = normalize_handle(handle)
    language = language_by_handle.get(key, {})
    return {
        "data_label": data_label,
        "handle": handle,
        "display_name": (following or {}).get("display_name", ""),
        "note": item.get("note", ""),
        "bio": (following or {}).get("bio", ""),
        "profile_url": (following or {}).get("profile_url") or f"https://x.com/{handle}",
        "in_following": in_following,
        "in_radar": True,
        "radar_status": item.get("status", ""),
        "language_label": language.get("language_label", ""),
        "source": "accounts_json",
        "fetched_at": fetched_at,
    }


def build_comparison(
    following_accounts: list[dict[str, Any]],
    radar_accounts: list[dict[str, Any]],
    *,
    language_by_handle: dict[str, dict[str, Any]] | None = None,
    fetched_at: str | None = None,
    username: str = "",
    expected_following_count: int | None = None,
    complete: bool | None = None,
) -> dict[str, Any]:
    fetched_at = fetched_at or datetime.now().astimezone().isoformat(timespec="seconds")
    language_by_handle = language_by_handle or load_language_by_handle(radar_accounts)

    following_by_key = {
        normalize_handle(item.get("handle", "")): item
        for item in following_accounts
        if normalize_handle(item.get("handle", ""))
    }
    active_radar = [item for item in radar_accounts if item.get("status") == "active"]
    radar_by_key = {
        normalize_handle(item.get("handle", "")): item
        for item in active_radar
        if normalize_handle(item.get("handle", ""))
    }

    following_rows = [
        _following_row(
            item,
            data_label="following",
            in_radar=key in radar_by_key,
            radar=radar_by_key.get(key),
            language_by_handle=language_by_handle,
            fetched_at=fetched_at,
        )
        for key, item in sorted(following_by_key.items())
    ]
    radar_rows = [
        _radar_row(
            item,
            data_label="radar",
            in_following=key in following_by_key,
            following=following_by_key.get(key),
            language_by_handle=language_by_handle,
            fetched_at=fetched_at,
        )
        for key, item in sorted(radar_by_key.items())
    ]
    gap1_rows = [
        _following_row(
            following_by_key[key],
            data_label="gap1_following_not_radar",
            in_radar=False,
            radar=None,
            language_by_handle=language_by_handle,
            fetched_at=fetched_at,
        )
        for key in sorted(set(following_by_key) - set(radar_by_key))
    ]
    gap2_rows = [
        _radar_row(
            radar_by_key[key],
            data_label="gap2_radar_not_following",
            in_following=False,
            following=None,
            language_by_handle=language_by_handle,
            fetched_at=fetched_at,
        )
        for key in sorted(set(radar_by_key) - set(following_by_key))
    ]

    collected_count = len(following_rows)
    if complete is None:
        complete = is_collection_complete(
            expected_following_count=expected_following_count,
            collected_count=collected_count,
        )
    status = completion_status(
        expected_following_count=expected_following_count,
        collected_count=collected_count,
    )

    return {
        "summary": {
            "username": username,
            "fetched_at": fetched_at,
            "expected_following_count": expected_following_count,
            "collected_following_count": collected_count,
            "completion_status": status,
            "following_count": len(following_rows),
            "radar_active_count": len(radar_rows),
            "gap1_count": len(gap1_rows),
            "gap2_count": len(gap2_rows),
            "complete": bool(complete),
        },
        "following": following_rows,
        "radar": radar_rows,
        "gap1": gap1_rows,
        "gap2": gap2_rows,
    }


def write_json_snapshot(snapshot: dict[str, Any], output_path: Path) -> Path:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(snapshot, ensure_ascii=False, indent=2), "utf-8")
    return output_path


def _style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="E2F0D9")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    ws.freeze_panes = "A2"
    for column in ws.columns:
        max_len = 0
        col = column[0].column
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 12), 48)


def _write_rows_sheet(wb, title: str, rows: list[dict[str, Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(ROW_COLUMNS)
    for row in rows:
        ws.append([row.get(column, "") for column in ROW_COLUMNS])
    _style_sheet(ws)


def write_xlsx_report(comparison: dict[str, Any], output_path: Path) -> Path:
    if Workbook is None:
        raise RuntimeError("缺少 openpyxl，请先安装依赖: pip install openpyxl")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"
    summary_ws.append(["metric", "value"])
    summary = comparison.get("summary", {})
    for field in SUMMARY_FIELDS:
        if field in summary:
            summary_ws.append([field, summary.get(field)])
    _style_sheet(summary_ws)

    for sheet_name, key in SHEET_MAP:
        _write_rows_sheet(wb, sheet_name, comparison.get(key, []))

    wb.save(output_path)
    return output_path


def _extract_visible_following_accounts(page: Page) -> list[dict[str, str]]:
    return page.evaluate(
        r"""() => {
            const bad = new Set([
                'home','explore','notifications','messages','i','settings',
                'compose','search','login','privacy','tos','jobs','help'
            ]);
            function cleanHandle(value) {
                return (value || '').replace(/^@/, '').trim();
            }
            function handleFromHref(href) {
                try {
                    const url = new URL(href, location.origin);
                    const parts = url.pathname.split('/').filter(Boolean);
                    if (parts.length !== 1) return '';
                    const candidate = parts[0];
                    if (bad.has(candidate.toLowerCase())) return '';
                    if (!/^[A-Za-z0-9_]{1,15}$/.test(candidate)) return '';
                    return candidate;
                } catch (_) {
                    return '';
                }
            }
            function usefulLine(line) {
                if (!line) return false;
                if (/^(Follow|Following|Follows you|Verified|Subscribe|Subscribed)$/i.test(line)) return false;
                return true;
            }
            const out = [];
            const cells = Array.from(document.querySelectorAll('[data-testid="UserCell"]'));
            for (const cell of cells) {
                const text = cell.innerText || '';
                const links = Array.from(cell.querySelectorAll('a[href]'));
                let handle = '';
                for (const link of links) {
                    handle = handleFromHref(link.getAttribute('href') || '');
                    if (handle) break;
                }
                const handleMatch = text.match(/@([A-Za-z0-9_]{1,15})/);
                if (handleMatch) handle = cleanHandle(handleMatch[1]) || handle;
                if (!handle) continue;

                const lines = text.split('\n').map(s => s.trim()).filter(usefulLine);
                const displayName = lines.find(line => !line.startsWith('@') && line !== handle) || handle;
                const bioLines = lines.filter(line => {
                    if (line === displayName) return false;
                    if (line === '@' + handle) return false;
                    if (line === handle) return false;
                    return !line.startsWith('@');
                });
                out.push({
                    handle,
                    display_name: displayName,
                    bio: bioLines.join(' ').trim(),
                    profile_url: `https://x.com/${handle}`
                });
            }
            return out;
        }"""
    )


def is_following_graphql_url(url: str) -> bool:
    try:
        parsed = urlparse(url or "")
    except Exception:
        return False
    path = (parsed.path or "").lower().rstrip("/")
    if "/graphql/" not in path:
        return False
    operation = path.split("/")[-1]
    return operation.startswith("following")


def fetch_following_accounts(
    page: Page,
    username: str,
    *,
    timeout: int = 30,
    limit: int | None = None,
    expected_count: int | None = None,
    idle_rounds: int = 15,
    scroll_pause: float = 1.2,
    max_scroll_rounds: int = 400,
) -> list[dict[str, str]]:
    username = display_handle(username)
    url = f"https://x.com/{username}/following"
    collected: dict[str, dict[str, str]] = {}
    network_collected: dict[str, dict[str, str]] = {}

    def handle_response(response) -> None:
        if not is_following_graphql_url(getattr(response, "url", "")):
            return
        try:
            payload = response.json()
        except Exception:
            return
        for row in extract_user_rows_from_graphql_payload(payload):
            key = normalize_handle(row.get("handle", ""))
            if key:
                network_collected[key] = merge_account_row(network_collected.get(key, {}), row)

    try:
        page.on("response", handle_response)
    except Exception:
        pass

    _log(f"打开 Following 页面: {url}")
    page.goto(url, wait_until="domcontentloaded", timeout=timeout * 1000)
    try:
        page.wait_for_selector('[data-testid="UserCell"]', timeout=min(timeout * 1000, 20000))
    except PlaywrightTimeoutError:
        _log("未等到 UserCell，继续尝试从当前页面提取", "WARN")

    unchanged = 0
    last_count = 0
    limit = limit or 0

    rounds = 0
    while rounds < max_scroll_rounds:
        rounds += 1
        for item in _extract_visible_following_accounts(page):
            key = normalize_handle(item.get("handle", ""))
            if key:
                collected[key] = merge_account_row(collected.get(key, {}), item)
        for key, item in network_collected.items():
            collected[key] = merge_account_row(collected.get(key, {}), item)

        if limit > 0 and len(collected) >= limit:
            break
        if expected_count and expected_count > 0 and len(collected) >= expected_count:
            break

        if len(collected) == last_count:
            unchanged += 1
        else:
            unchanged = 0
            last_count = len(collected)

        if unchanged >= idle_rounds:
            break

        try:
            page.mouse.wheel(0, 2400)
        except Exception:
            page.evaluate("window.scrollBy(0, Math.max(window.innerHeight * 1.5, 1600))")
        try:
            page.wait_for_timeout(int(scroll_pause * 1000))
        except Exception:
            time.sleep(scroll_pause)

    try:
        if hasattr(page, "remove_listener"):
            page.remove_listener("response", handle_response)
        elif hasattr(page, "off"):
            page.off("response", handle_response)
    except Exception:
        pass

    rows = list(collected.values())
    rows.sort(key=lambda item: normalize_handle(item.get("handle", "")))
    return rows[:limit] if limit > 0 else rows


def _expected_following_count(page: Page, username: str, timeout: int) -> int:
    try:
        from fetch_follower_stats import fetch_profile_stats

        stats = fetch_profile_stats(page, username, timeout=timeout)
        return int(stats.get("following", -1))
    except Exception as exc:
        _log(f"读取 Following 数量失败: {exc}", "WARN")
        return -1


def run_sync(args) -> dict[str, Any]:
    username = display_handle(args.username)
    fetched_at = datetime.now().astimezone().isoformat(timespec="seconds")
    date_stamp = datetime.now().strftime("%Y-%m-%d")
    json_output = Path(args.json_output) if args.json_output else FOLLOWING_DIR / f"{date_stamp}.json"
    latest_output = FOLLOWING_DIR / "myfollowing_latest.json"
    xlsx_output = Path(args.xlsx_output) if args.xlsx_output else DEFAULT_XLSX
    profile_dir = Path(args.profile_dir) if args.profile_dir else DEFAULT_PROFILE_DIR

    session = None
    try:
        profile_dir.mkdir(parents=True, exist_ok=True)
        session = create_session(profile_dir=profile_dir, headless=not args.no_headless)
        page = session.create_task_page()
        session.page = page

        if not session.check_login_status(timeout=args.timeout * 1000):
            raise RuntimeError("X 登录态不可用，请先运行: python xpost.py post-login")

        expected_count = _expected_following_count(page, username, args.timeout)
        following = fetch_following_accounts(
            page,
            username,
            timeout=args.timeout,
            limit=args.limit,
            expected_count=expected_count if expected_count >= 0 else None,
            idle_rounds=args.idle_rounds,
        )
    finally:
        if session:
            session.close()

    radar_accounts = load_radar_accounts(ACCOUNTS_JSON)
    effective_expected = expected_count if expected_count >= 0 else None
    status = completion_status(
        expected_following_count=effective_expected,
        collected_count=len(following),
        limit=args.limit,
    )
    complete = status == "complete"

    comparison = build_comparison(
        following,
        radar_accounts,
        fetched_at=fetched_at,
        username=username,
        expected_following_count=effective_expected,
        complete=complete,
    )
    comparison["summary"]["json_snapshot"] = str(json_output)
    comparison["summary"]["xlsx_report"] = str(xlsx_output)
    comparison["summary"]["completion_status"] = status

    snapshot = {
        "ok": True,
        "username": username,
        "fetched_at": fetched_at,
        "expected_following_count": effective_expected,
        "collected_following_count": len(following),
        "complete": complete,
        "completion_status": status,
        "following": following,
        "comparison": comparison,
    }
    write_json_snapshot(snapshot, json_output)
    write_json_snapshot(snapshot, latest_output)
    write_xlsx_report(comparison, xlsx_output)

    return {
        "ok": True,
        "username": username,
        "json_snapshot": str(json_output),
        "latest_snapshot": str(latest_output),
        "xlsx_report": str(xlsx_output),
        "summary": comparison["summary"],
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Fetch X following list and compare with radar accounts")
    parser.add_argument("username", help="X username without @")
    parser.add_argument("--limit", type=int, default=0, help="Max following accounts to collect; 0 means unlimited")
    parser.add_argument("--timeout", type=int, default=30, help="Page load timeout seconds")
    parser.add_argument("--idle-rounds", type=int, default=15, help="Stop after N scroll rounds without new accounts")
    parser.add_argument("--profile-dir", help="Chrome profile directory")
    parser.add_argument("--json-output", help="JSON snapshot path")
    parser.add_argument("--xlsx-output", help="XLSX report path")
    parser.add_argument("--no-headless", action="store_true", help="Use visible Chrome window")
    parser.add_argument("--json-only", action="store_true", help="Only print JSON result")
    args = parser.parse_args()

    try:
        result = run_sync(args)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result.get("ok") else 1
    except KeyboardInterrupt:
        return 130
    except Exception as exc:
        result = {"ok": False, "error": str(exc)}
        print(json.dumps(result, ensure_ascii=False, indent=2))
        if not args.json_only:
            _log(f"错误: {exc}", "ERROR")
        return 1


if __name__ == "__main__":
    sys.exit(main())
