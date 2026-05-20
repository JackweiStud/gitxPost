from pathlib import Path

import openpyxl

import following_sync


def test_graphql_payload_extracts_followers_count():
    payload = {
        "data": {
            "user": {
                "result": {
                    "timeline": {
                        "timeline": {
                            "instructions": [
                                {
                                    "entries": [
                                        {
                                            "content": {
                                                "itemContent": {
                                                    "user_results": {
                                                        "result": {
                                                            "legacy": {
                                                                "screen_name": "0xFortuneRust",
                                                                "name": "FortuneRust",
                                                                "description": "WebSite, Mobile, Game, Blockchain developer",
                                                                "followers_count": 3129,
                                                            }
                                                        }
                                                    }
                                                }
                                            }
                                        }
                                    ]
                                }
                            ]
                        }
                    }
                }
            }
        }
    }

    rows = following_sync.extract_user_rows_from_graphql_payload(payload)

    assert rows[0]["handle"] == "0xFortuneRust"
    assert rows[0]["followers_count"] == 3129


def test_graphql_payload_extracts_handle_from_core_and_followers_from_legacy():
    payload = {
        "data": {
            "user": {
                "result": {
                    "timeline": {
                        "timeline": {
                            "instructions": [
                                {
                                    "entries": [
                                        {
                                            "content": {
                                                "itemContent": {
                                                    "user_results": {
                                                        "result": {
                                                            "core": {
                                                                "screen_name": "0xFortuneRust",
                                                                "name": "FortuneRust",
                                                            },
                                                            "legacy": {
                                                                "description": "WebSite, Mobile, Game, Blockchain developer",
                                                                "followers_count": 3129,
                                                            },
                                                        }
                                                    }
                                                }
                                            }
                                        }
                                    ]
                                }
                            ]
                        }
                    }
                }
            }
        }
    }

    rows = following_sync.extract_user_rows_from_graphql_payload(payload)

    assert rows[0]["handle"] == "0xFortuneRust"
    assert rows[0]["display_name"] == "FortuneRust"
    assert rows[0]["followers_count"] == 3129


def test_xlsx_following_sheet_includes_followers_count(tmp_path: Path):
    comparison = following_sync.build_comparison(
        [
            {
                "handle": "0xFortuneRust",
                "display_name": "FortuneRust",
                "bio": "WebSite, Mobile, Game, Blockchain developer",
                "profile_url": "https://x.com/0xFortuneRust",
                "followers_count": 3129,
            }
        ],
        [],
        fetched_at="2026-05-20T12:00:00+08:00",
        expected_following_count=1,
    )
    out = tmp_path / "myfollowing.xlsx"

    following_sync.write_xlsx_report(comparison, out)

    wb = openpyxl.load_workbook(out)
    ws = wb["Following"]
    headers = [cell.value for cell in ws[1]]
    assert "followers_count" in headers
    followers_col = headers.index("followers_count") + 1
    assert ws.cell(row=2, column=followers_col).value == 3129


def test_fill_missing_followers_counts_uses_profile_stats_fetcher():
    rows = [
        {"handle": "with_count", "followers_count": 10},
        {"handle": "missing_count"},
    ]

    def fake_fetcher(_page, username, timeout):
        assert username == "missing_count"
        assert timeout == 30
        return {"followers": 3129}

    updated = following_sync.fill_missing_followers_counts(
        page=object(),
        rows=rows,
        timeout=30,
        fetcher=fake_fetcher,
    )

    assert updated == 1
    assert rows[0]["followers_count"] == 10
    assert rows[1]["followers_count"] == 3129
